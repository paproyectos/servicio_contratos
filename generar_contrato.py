#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de contratos Word desde plantilla con marcadores:
  - Acepta: [campo], {{campo}}, <<campo>> (tolerante a espacios)
  - Reemplazo UNIVERSAL en todos los w:t (cuerpo, tablas, header, footer, textboxes)
  - Canoniza marcadores partidos dentro de un mismo párrafo/celda/header/footer
  - Expande repeat: [repeat:items], {{repeat:items}} o <<repeat:items>>
    (acepta también la variante con tilde: ítems)
  - Si no envías "items", pero sí item_descripcion_1..N, los convierte automáticamente
  - Valida campos, cuotas, sumas y fechas
  - Formatea CLP con miles y sufijo ".-"
  - Opciones: --listar-marcadores y --resaltar-pendientes

Uso:
  python generar_contrato.py --datos datos.json --plantilla plantilla.docx --salida outdir/

Autoría del método y del script:
  - Asistente: ChatGPT (GPT-5 Thinking)
  - Colaborador humano: Pa ♥

Créditos: desarrollado colaborativamente en 2025 con soporte de ChatGPT.
Licencia sugerida: MIT (puedes cambiarla).

"""
import argparse
import json
import os
import re
import unicodedata
from datetime import datetime
from typing import Dict, Any, List, Tuple

# IO opcional de .xlsx/.csv
try:
    import pandas as pd
except Exception:
    pd = None

from docx import Document
from docx.text.paragraph import Paragraph

from copy import deepcopy

# ------------------------- Configuración delimitadores ------------------------
# Aceptamos estos tres delimitadores simultáneamente (tolerante a espacios):
DELIMS = [
    ("[", "]"),
    ("{{", "}}"),
    ("<<", ">>"),
]

# ------------------------- Utilidades texto/fechas ----------------------------
MESES_ES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
    "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
    "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12
}

def quitar_acentos(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"
    )

def parse_fecha_es(fecha: str) -> datetime:
    """
    Admite:
      '01 de agosto de 2025', '1 agosto 2025', '2025-08-01', '01/08/2025'
    """
    fecha = (fecha or "").strip()
    if not fecha:
        raise ValueError("Fecha vacía.")
    # ISO
    try:
        return datetime.fromisoformat(fecha)
    except Exception:
        pass
    # dd/mm/yyyy
    m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$', fecha)
    if m:
        d, mth, y = map(int, m.groups())
        return datetime(y, mth, d)
    # '1 de agosto de 2025' o '1 agosto 2025'
    m = re.match(r'^(\d{1,2})\s*(?:de\s*)?([A-Za-záéíóúñÑ]+)\s*(?:de\s*)?(\d{4})$', fecha, flags=re.IGNORECASE)
    if m:
        d = int(m.group(1))
        mes = quitar_acentos(m.group(2).lower())
        y = int(m.group(3))
        if mes in MESES_ES:
            return datetime(y, MESES_ES[mes], d)
    raise ValueError(f"Fecha no reconocida: '{fecha}'")

def clp_formato(n: int | str) -> str:
    """1980000 -> '1.980.000.-' (sin $). Si no hay dígitos, '--'."""
    if isinstance(n, str):
        digits = re.sub(r'[^\d]', '', n)
        if digits == '':
            return '--'
        n = int(digits)
    s = f"{n:,}".replace(",", ".")
    return f"{s}.-"

def partir_nombre(nombre: str) -> Tuple[str, str]:
    toks = [t for t in (nombre or "").strip().split() if t]
    if not toks:
        return ("", "")
    nombre_pila = toks[0]
    if len(toks) >= 3:
        apellido = toks[-2]
    elif len(toks) == 2:
        apellido = toks[-1]
    else:
        apellido = toks[0]
    return (nombre_pila, apellido)

# ------------------------- Carga de datos -------------------------------------
def cargar_datos(path: str) -> Dict[str, Any]:
    ext = os.path.splitext(path.lower())[1]
    if ext == ".json":
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    elif ext in (".xlsx", ".xls"):
        if pd is None:
            raise RuntimeError("pandas no disponible para leer Excel.")
        df = pd.read_excel(path)
        if df.shape[0] != 1:
            raise ValueError("El Excel debe contener exactamente 1 fila con los datos.")
        return df.iloc[0].to_dict()
    elif ext == ".csv":
        if pd is None:
            # CSV simple 'clave,valor' -> dict
            datos: Dict[str, str] = {}
            with open(path, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if not line or line.startswith("#"):
                        continue
                    if "," in line:
                        k, v = line.split(",", 1)
                        datos[k.strip()] = v.strip()
            return datos
        else:
            df = pd.read_csv(path)
            if df.shape[0] != 1:
                raise ValueError("El CSV debe contener exactamente 1 fila con los datos.")
            return df.iloc[0].to_dict()
    else:
        raise ValueError(f"Formato de datos no soportado: {ext}")

# ------------------------- Extracción de marcadores (universal) ---------------
def _tokens_from_text(txt: str) -> List[str]:
    if not txt:
        return []
    tokens = set()
    patrones = [
        re.compile(r'\[([^\]\r\n]+)\]'),
        re.compile(r'\{\{\s*([^}\r\n]+?)\s*\}\}'),
        re.compile(r'\<\<\s*([^>\r\n]+?)\s*\>\>')
    ]
    for pat in patrones:
        for m in pat.finditer(txt):
            token = (m.group(1) or "").strip()
            if token.startswith("repeat:") or token.startswith("/repeat"):
                continue
            tokens.add(token)
    return list(tokens)

def extraer_marcadores_universal(doc: Document) -> List[str]:
    """Lee todos los w:t de cuerpo + headers/footers y extrae tokens."""
    tokens = set()
    parts = [doc.part]
    for section in doc.sections:
        if section.header: parts.append(section.header.part)
        if section.footer: parts.append(section.footer.part)
    for part in parts:
        for t in part.element.xpath('.//w:t'):
            for tok in _tokens_from_text(t.text or ''):
                tokens.add(tok)
    return sorted(tokens)

# ------------------------- Validaciones ---------------------------------------
def validar(datos: Dict[str, Any], doc: Document):
    """
    - Todos los tokens detectados deben tener valor NO vacío,
      excepto: 'item_descripcion' y 'mes_i'/'monto_i' (permitidos como '--').
    - numero_cuotas debe coincidir con la cantidad de pares (mes_i, monto_i) no vacíos.
    - La suma de montos de cuotas no vacías debe ser igual a 'monto'.
    - fecha_inicio ≤ fecha_termino.
    """
    marcadores = extraer_marcadores_universal(doc)
    faltantes = []
    for m in marcadores:
        val = str(datos.get(m, "")).strip()
        if m == 'item_descripcion' or re.match(r'^(mes|monto)_\d+$', m):
            continue
        if val == "":
            faltantes.append(m)
    if faltantes:
        raise ValueError("Faltan datos obligatorios para: " + ", ".join(faltantes))

    # numero_cuotas vs pares (mes_i, monto_i)
    try:
        num = int(re.sub(r'[^\d]', '', str(datos.get("numero_cuotas", "0"))))
    except Exception:
        raise ValueError("numero_cuotas debe ser un entero.")

    pares_no_vacios = 0
    for i in range(1, 13):
        mi = str(datos.get(f"mes_{i}", "")).strip()
        mo = str(datos.get(f"monto_{i}", "")).strip()
        if mi not in ('', '--') and mo not in ('', '--'):
            pares_no_vacios += 1
    if pares_no_vacios != num:
        raise ValueError(f"numero_cuotas={num} no coincide con pares no vacíos={pares_no_vacios}.")

    # Σ monto_i == monto (considerando solo cuotas no vacías)
    total = 0
    for i in range(1, 13):
        mi = str(datos.get(f"mes_{i}", "")).strip()
        mo = str(datos.get(f"monto_{i}", "")).strip()
        if mi not in ('', '--') and mo not in ('', '--'):
            digits = re.sub(r'[^\d]', '', mo)
            if digits:
                total += int(digits)
    monto_total = int(re.sub(r'[^\d]', '', str(datos.get("monto", "0")) or "0"))
    if total != monto_total:
        raise ValueError(f"La suma de cuotas ({total}) difiere del monto total ({monto_total}).")

    # fechas: inicio ≤ término
    fi = parse_fecha_es(str(datos.get("fecha_inicio", "")).strip())
    ft = parse_fecha_es(str(datos.get("fecha_termino", "")).strip())
    if fi > ft:
        raise ValueError("fecha_inicio es posterior a fecha_termino.")

# ------------------------- Canonización por párrafo (preserva formato) --------
def canonizar_placeholders_en_parrafos_de_part(part, delims):
    """
    Recorre todos los w:p de una 'part' (documento, header o footer).
    Si un marcador {{...}} / <<...>> / [...] está cortado en varios w:t,
    lo une y coloca el texto final en el w:t cuyo run tenga el "mejor" formato
    (b/i/sz), vaciando los demás. Así preserva negrita, cursiva y tamaño.
    """
    def run_score_for_t(t):
        # Puntúa el estilo del run padre: b=5, i=4, tamaño explícito=3
        score = 0
        r = t.getparent()  # w:r
        if r is None:
            return score
        rPr = r.xpath('w:rPr')
        if not rPr:
            return score
        rPr = rPr[0]
        if rPr.xpath('w:b'):
            score += 5
        if rPr.xpath('w:i'):
            score += 4
        if rPr.xpath('w:sz'):
            score += 3
        return score

    for p in part.element.xpath('.//w:p'):
        ts = p.xpath('.//w:t')
        if not ts:
            continue

        buffering = False
        buf_text = ""
        buf_nodes = []
        open_delim = None  # (ldel, rdel)

        i = 0
        while i < len(ts):
            t = ts[i]
            text = t.text or ""
            j = 0
            while j < len(text):
                if not buffering:
                    started = False
                    for ldel, rdel in delims:
                        L = len(ldel)
                        if text[j:j+L] == ldel:
                            buffering = True
                            open_delim = (ldel, rdel)
                            buf_text = ldel
                            buf_nodes = [t]  # empezamos buffer en este t
                            j += L
                            started = True
                            break
                    if not started:
                        j += 1
                else:
                    buf_text += text[j]
                    if not buf_nodes or buf_nodes[-1] is not t:
                        buf_nodes.append(t)
                    # ¿cerró?
                    ldel, rdel = open_delim
                    if buf_text.endswith(rdel):
                        # Elegir el w:t "mejor formateado" para alojar el marcador
                        best_t = max(buf_nodes, key=run_score_for_t)
                        best_t.text = buf_text
                        # Vaciar el resto de nodos que participaron
                        for nn in buf_nodes:
                            if nn is not best_t:
                                nn.text = ""
                        # Reset
                        buffering = False
                        buf_text = ""
                        buf_nodes = []
                        open_delim = None
                    j += 1
            # Si el marcador sigue abierto, y este t no es el primero del buffer, vaciarlo
            if buffering and t is not (buf_nodes[0] if buf_nodes else None):
                t.text = ""
            i += 1

def canonizar_en_todo_el_doc(doc: Document):
    parts = [doc.part]
    for section in doc.sections:
        if section.header: parts.append(section.header.part)
        if section.footer: parts.append(section.footer.part)
    for part in parts:
        canonizar_placeholders_en_parrafos_de_part(part, DELIMS)

# ------------------------- Reemplazo UNIVERSAL --------------------------------
def reemplazar_en_todos_los_textos(doc: Document, mapping: Dict[str, str]):
    """
    Reemplaza en TODOS los w:t de todas las parts (cuerpo, tablas, encabezados, pies,
    cuadros de texto y SDTs). Tolerante a espacios y a múltiples delimitadores.
    """
    parts = [doc.part]
    for section in doc.sections:
        if section.header: parts.append(section.header.part)
        if section.footer: parts.append(section.footer.part)

    for part in parts:
        for t in part.element.xpath('.//w:t'):
            txt = t.text or ''
            for k, v in mapping.items():
                for ldel, rdel in DELIMS:
                    pat = re.compile(re.escape(ldel) + r"\s*" + re.escape(k) + r"\s*" + re.escape(rdel))
                    txt = pat.sub(v, txt)
            t.text = txt

# ------------------------- Resaltado de pendientes (debug) --------------------
def resaltar_pendientes_universal(doc: Document):
    """Marca con '¡¡ ... !!' cualquier token no reemplazado en todos los w:t."""
    rx = re.compile(r'(\{\{\s*[^}\r\n]+\s*\}\}|\<\<\s*[^>\r\n]+\s*\>\>|\[[^\]\r\n]+\])')
    parts = [doc.part]
    for section in doc.sections:
        if section.header: parts.append(section.header.part)
        if section.footer: parts.append(section.footer.part)
    for part in parts:
        for t in part.element.xpath('.//w:t'):
            txt = t.text or ''
            if rx.search(txt):
                t.text = rx.sub(lambda m: f"¡¡{m.group(1)}!!", txt)

# ------------------------- Bloques repetibles (robusto, preserva formato) -----
def expandir_repeat_items(doc: Document, items: list[dict[str, str]]):
    """
    Expande el bloque repeat en el MISMO lugar (funciona en párrafos de tablas/txbx),
    clonando el párrafo plantilla y REEMPLAZANDO item_descripcion aunque el marcador
    esté cortado en múltiples w:t. Conserva el formato del run mejor formateado (b/i/sz).
    Admite: [repeat:items] ... [/repeat:items], {{...}}, <<...>> (también 'ítems').
    """
    root = doc.part.element

    # --- helpers -------------------------------------------------------------
    def p_text(p):
        return "".join(t.text or "" for t in p.xpath('.//w:t'))

    OPEN_RX  = re.compile(
        r'^\s*(?:\[\s*repeat\s*:\s*(?:items|ítems)\s*\]'
        r'|\{\{\s*repeat\s*:\s*(?:items|ítems)\s*\}\}'
        r'|\<\<\s*repeat\s*:\s*(?:items|ítems)\s*\>\>)\s*$', re.IGNORECASE)

    CLOSE_RX = re.compile(
        r'^\s*(?:\[\s*/\s*repeat\s*:\s*(?:items|ítems)\s*\]'
        r'|\{\{\s*/\s*repeat\s*:\s*(?:items|ítems)\s*\}\}'
        r'|\<\<\s*/\s*repeat\s*:\s*(?:items|ítems)\s*\>\>)\s*$', re.IGNORECASE)

    # acepta [item_descripcion], {{ item_descripcion }}, << item_descripcion >>
    ITEM_OPEN = ( "[", "{{", "<<" )
    ITEM_CLOSE= ( "]", "}}", ">>" )
    ITEM_NAME_RX = re.compile(r'^\s*item_descripcion\s*$', re.IGNORECASE)

    def run_score_for_t(t):
        # puntúa run padre: b=5, i=4, w:sz=3 (para elegir el destino del texto)
        score = 0
        r = t.getparent()
        if r is None: return score
        rPrs = r.xpath('w:rPr')
        if not rPrs: return score
        rPr = rPrs[0]
        if rPr.xpath('w:b'):  score += 5
        if rPr.xpath('w:i'):  score += 4
        if rPr.xpath('w:sz'): score += 3
        return score

    def replace_item_placeholder_in_clone(clone_p, replacement_text: str):
        """
        Reemplaza el marcador item_descripcion aunque esté partido en varios w:t.
        Coloca el reemplazo en el w:t cuyo run tiene mejor formato y vacía el resto.
        Si hay múltiples apariciones, reemplaza todas.
        """
        ts = clone_p.xpath('.//w:t')
        if not ts:
            return

        # state machine que cruza múltiples w:t
        buffering = False
        buf_nodes = []
        buf_text = ""
        cur_delims = None  # (ldel, rdel)

        def maybe_flush_if_item():
            # Si buf_text contiene {{ item_descripcion }}, hacer el reemplazo
            nonlocal buf_nodes, buf_text, buffering, cur_delims
            # Extraer el contenido sin delimitadores
            for ldel, rdel in zip(ITEM_OPEN, ITEM_CLOSE):
                if buf_text.startswith(ldel) and buf_text.endswith(rdel):
                    inner = buf_text[len(ldel):-len(rdel)]
                    if ITEM_NAME_RX.match(inner):
                        # elegir mejor w:t para alojar el texto final
                        best_t = max(buf_nodes, key=run_score_for_t)
                        best_t.text = replacement_text
                        for nn in buf_nodes:
                            if nn is not best_t:
                                nn.text = ""
                        # reset
                        buffering = False
                        buf_nodes = []
                        buf_text = ""
                        cur_delims = None
                        return True
            return False

        for t in ts:
            text = t.text or ""
            i = 0
            while i < len(text):
                if not buffering:
                    # ¿empieza algún delimitador aquí?
                    started = False
                    for ldel in ITEM_OPEN:
                        L = len(ldel)
                        if text[i:i+L] == ldel:
                            buffering = True
                            cur_delims = (ldel, ITEM_CLOSE[ITEM_OPEN.index(ldel)])
                            buf_nodes = [t]
                            buf_text = ldel
                            i += L
                            started = True
                            break
                    if not started:
                        i += 1
                else:
                    buf_text += text[i]
                    if t not in buf_nodes:
                        buf_nodes.append(t)
                    # ¿cerró?
                    ldel, rdel = cur_delims
                    R = len(rdel)
                    if buf_text.endswith(rdel):
                        # intentamos reemplazar si es item_descripcion
                        replaced = maybe_flush_if_item()
                        if not replaced:
                            # no era item_descripcion → no tocar, reset sin cambiar textos
                            buffering = False
                            buf_nodes = []
                            buf_text = ""
                            cur_delims = None
                    i += 1
            # si seguimos dentro de un marcador, vaciamos este t (para no duplicar)
            if buffering and t is not buf_nodes[0]:
                t.text = ""

        # Si terminara con marcador sin cerrar, lo dejamos tal cual

    # --- 1) localizar bloque en TODOS los párrafos XML -----------------------
    ps = root.xpath('.//w:p')
    start = end = None
    for i, p in enumerate(ps):
        if OPEN_RX.match(p_text(p)):
            start = i
            break
    if start is None:
        return  # no hay bloque

    for j in range(start + 1, len(ps)):
        if CLOSE_RX.match(p_text(ps[j])):
            end = j
            break
    if end is None or end <= start:
        return  # bloque mal formado

    # --- 2) párrafo plantilla (el que contiene item_descripcion) -------------
    template_p = None
    for k in range(start + 1, end):
        txt = p_text(ps[k]).lower()
        if ("item" in txt and "descripcion" in txt) or "item_descripcion" in txt:
            template_p = ps[k]
            break
    if template_p is None:
        return

    close_p = ps[end]

    # --- 3) insertar clones antes del cierre y reemplazar dentro del clon ----
    for it in items or []:
        clone = deepcopy(template_p)
        replace_item_placeholder_in_clone(clone, str(it.get("item_descripcion", "") or ""))
        close_p.addprevious(clone)

    # --- 4) eliminar apertura..cierre (incluye la plantilla original) --------
    for idx in range(end, start - 1, -1):
        p = ps[idx]
        par = p.getparent()
        if par is not None:
            par.remove(p)

# ------------------------- Mapping y normalización ----------------------------
def build_mapping(datos: Dict[str, Any]) -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    for k, v in datos.items():
        mapping[k] = "" if v is None else str(v)

    if "monto" in mapping:
        mapping["monto"] = clp_formato(mapping["monto"])

    for i in range(1, 13):
        key_monto = f"monto_{i}"
        key_mes = f"mes_{i}"
        vi = (mapping.get(key_monto, "") or "").strip()
        mi = (mapping.get(key_mes, "") or "").strip()
        mapping[key_monto] = clp_formato(vi) if vi not in ("", "--") else "--"
        mapping[key_mes] = mi if mi not in ("",) else "--"

    return mapping

# ------------------------- Proceso principal ----------------------------------
def generar_contrato(datos_path: str, plantilla_path: str, salida_dir: str,
                     listar_marcadores: bool = False,
                     resaltar: bool = False) -> str | None:
    # Cargar datos y documento
    datos = cargar_datos(datos_path)
    doc = Document(plantilla_path)

    if listar_marcadores:
        print("Marcadores encontrados:")
        for t in extraer_marcadores_universal(doc):
            print("-", t)
        return None

    # Construir items si vinieron como item_descripcion_1..N
    items = datos.get("items", [])
    if (not items) and any(k.startswith("item_descripcion_") for k in datos.keys()):
        ordenados = sorted(
            [(int(re.sub(r'\D', '', k) or 0), v) for k, v in datos.items() if k.startswith("item_descripcion_")],
            key=lambda x: x[0]
        )
        items = [{"item_descripcion": v} for _, v in ordenados]

    # Validar vs plantilla
    validar(datos, doc)

    # Mapping (formateos y '--')
    mapping = build_mapping(datos)

    # Expandir items si hay
    if isinstance(items, list) and items:
        expandir_repeat_items(doc, items)

    # Canonizar marcadores partidos dentro de cada párrafo/celda/header/footer (preserva formato)
    canonizar_en_todo_el_doc(doc)

    # REEMPLAZO UNIVERSAL (una pasada en todos los w:t de todas las parts)
    reemplazar_en_todos_los_textos(doc, mapping)

    # Depuración visual (opcional)
    if resaltar:
        resaltar_pendientes_universal(doc)

    # Nombre de archivo
    nombre_pt = str(datos.get("nombre_PT", "")).strip()
    nombre, apellido = partir_nombre(nombre_pt)
    apellido = quitar_acentos(apellido).title()
    nombre_simple = quitar_acentos(nombre).title()
    codigo = str(datos.get("codigo", "")).strip()

    etapa = str(datos.get("etapa", "")).strip()
    m = re.search(r'(\d{4})', etapa or "")
    if m:
        etapa_aaaa = m.group(1)
    else:
        try:
            etapa_aaaa = str(parse_fecha_es(str(datos.get("fecha_inicio", "")).strip()).year)
        except Exception:
            etapa_aaaa = "XXXX"

    nombre_archivo = f"{apellido}, {nombre_simple}, Contrato{codigo}Etapa{etapa_aaaa}.docx"

    # Metadatos del documento (autoría)
    try:
        cp = doc.core_properties
        cp.author = (cp.author or "Pa")
        cp.last_modified_by = "Pa"
        comment_author = __author__ if "__author__" in globals() else "ChatGPT (GPT-5 Thinking) + Usuario"
        version_str   = __version__ if "__version__" in globals() else "1.0.0"
        cp.comments = f"Generado con script de Pa con soporte de {comment_author}, v{version_str}."
        if not cp.title:
            cp.title = "Contrato"
        if not cp.subject:
            cp.subject = "Generado desde plantilla con marcadores"
    except Exception:
        # No bloquear la generación si no se pueden escribir metadatos
        pass

    # Guardar
    os.makedirs(salida_dir, exist_ok=True)
    out_path = os.path.join(salida_dir, nombre_archivo)
    doc.save(out_path)
    return out_path

def main():
    ap = argparse.ArgumentParser(description="Generador de contratos .docx desde plantilla con marcadores.")
    ap.add_argument("--datos", required=False, help="Ruta a datos (.json, .csv, .xlsx) con encabezados = marcadores.")
    ap.add_argument("--plantilla", required=False, help="Ruta a plantilla .docx (admite [], {{}}, <<>>).")
    ap.add_argument("--salida", required=False, help="Directorio de salida.")
    ap.add_argument("--listar-marcadores", action="store_true",
                    help="Muestra los marcadores detectados en la plantilla y sale.")
    ap.add_argument("--resaltar-pendientes", action="store_true",
                    help="Resalta marcadores NO reemplazados encerrándolos en '¡¡ ... !!'.")
    ap.add_argument("--about", action="store_true",
                    help="Muestra información de autoría y versión y sale.")
    args = ap.parse_args()

    # Mostrar autoría y salir
    if args.about:
        try:
            print(f"Generar contrato - versión {__version__}")
        except NameError:
            print("Generar contrato - versión (no definida)")
        try:
            print(f"Autoría: {__author__}")
        except NameError:
            print("Autoría: ChatGPT (GPT-5 Thinking) + Usuario")
        print("Créditos: desarrollado colaborativamente en 2025 con soporte de ChatGPT (GPT-5 Thinking).")
        return

    if args.listar_marcadores:
        if not args.plantilla:
            raise SystemExit("Para --listar-marcadores debes indicar --plantilla.")
        doc = Document(args.plantilla)
        print("Marcadores encontrados:")
        for t in extraer_marcadores_universal(doc):
            print("-", t)
        return

    if not (args.datos and args.plantilla and args.salida):
        raise SystemExit("Debes indicar --datos, --plantilla y --salida (o usa --listar-marcadores / --about).")

    out = generar_contrato(args.datos, args.plantilla, args.salida,
                           listar_marcadores=False,
                           resaltar=args.resaltar_pendientes)
    if out:
        print(out)

if __name__ == "__main__":
    main()